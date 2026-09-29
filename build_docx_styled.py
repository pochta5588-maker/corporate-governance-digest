#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Сборка Дайджеста в акцентном оформлении (оформленный вариант).

Палитра:
- оранжевый #ff6d22 (акцент: линии-подчёркивания, маркеры новостей)
- тёмно-серый #565a5c (заголовки)
- серый #565656 (текст)
- светлый #979797 (служебные подписи: колонтитулы, примечания)
"""
import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# --- палитра ---
ACCENT = 'FF6D22'
DARKGRAY = '565A5C'
GRAY = '565656'
LIGHTGRAY = '979797'
LINKBLUE = '0563C1'

def rgb(hexstr):
    return RGBColor.from_string(hexstr)

def set_font(run, size=12, bold=False, italic=False, color=None, name='Times New Roman'):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = rgb(color)

def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Times New Roman'); rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rPr.append(rFonts)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '24')  # 12pt
    rPr.append(sz)
    color = OxmlElement('w:color'); color.set(qn('w:val'), LINKBLUE); rPr.append(color)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    new_run.append(rPr)
    t = OxmlElement('w:t'); t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

LINK_RE = re.compile(r'\[([^\]]+)\]\(([^)]+)\)')
BOLD_RE = re.compile(r'\*\*([^*]+)\*\*')

def _add_bold_segments(par, text, size, base_bold, base_color, font='Times New Roman'):
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            r = par.add_run(text[pos:m.start()]); set_font(r, size, base_bold, color=base_color, name=font)
        r = par.add_run(m.group(1)); set_font(r, size, True, color=base_color, name=font)
        pos = m.end()
    if pos < len(text):
        r = par.add_run(text[pos:]); set_font(r, size, base_bold, color=base_color, name=font)

def add_runs_with_links(par, text, base_size=12, base_bold=False, base_color=GRAY, font='Times New Roman'):
    pos = 0
    for m in LINK_RE.finditer(text):
        seg = text[pos:m.start()]
        _add_bold_segments(par, seg, base_size, base_bold, base_color, font)
        add_hyperlink(par, m.group(1), m.group(2))
        pos = m.end()
    _add_bold_segments(par, text[pos:], base_size, base_bold, base_color, font)

def para(doc, text, size=12, bold=False, italic=False, space_after=6, align=None, color=GRAY, font='Times New Roman'):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.15
    if align is not None: p.alignment = align
    add_runs_with_links(p, text, size, bold, color, font)
    if italic:
        for r in p.runs:
            r.italic = True
    return p

def set_bottom_border(par, color=ACCENT, sz=16, space=6):
    """Линия под абзацем (sz в восьмых пункта)."""
    pPr = par._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(sz))
    bottom.set(qn('w:space'), str(space))
    bottom.set(qn('w:color'), color)
    pbdr.append(bottom)
    pPr.append(pbdr)

def add_page_number_footer(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # линия сверху + номер страницы
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    top = OxmlElement('w:top')
    top.set(qn('w:val'), 'single'); top.set(qn('w:sz'), '6')
    top.set(qn('w:space'), '4'); top.set(qn('w:color'), ACCENT)
    pbdr.append(top); pPr.append(pbdr)
    r1 = p.add_run('Дайджест новостей по корпоративному управлению · стр. ')
    set_font(r1, 9, color=LIGHTGRAY, name='Arial')
    r = p.add_run()
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    r._r.append(fldChar1); r._r.append(instrText); r._r.append(fldChar2)
    set_font(r, 9, color=LIGHTGRAY, name='Arial')

def add_text_header(doc):
    header = doc.sections[0].header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run('Дайджест по корпоративному управлению · квартальная сводка')
    set_font(r, 9, color=LIGHTGRAY, name='Arial')
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '4'); bottom.set(qn('w:color'), LIGHTGRAY)
    pbdr.append(bottom); pPr.append(pbdr)

def build(src, dst):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2); section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5); section.right_margin = Cm(1.5)

    add_text_header(doc)
    add_page_number_footer(doc)

    with open(src, encoding='utf-8') as f:
        lines = f.read().splitlines()

    for line in lines:
        if not line.strip():
            continue
        if line.startswith('# '):
            # титул
            para(doc, line[2:].strip(), size=18, bold=True, space_after=8,
                 align=WD_ALIGN_PARAGRAPH.LEFT, color=DARKGRAY, font='Arial')
        elif line.startswith('Период:'):
            p = para(doc, line.strip(), size=11, space_after=2, color=DARKGRAY, font='Arial')
            set_bottom_border(p, ACCENT, sz=18, space=8)  # линия под титульным блоком
        elif line.startswith('Примечание:'):
            para(doc, line.strip(), size=10, italic=True, space_after=14, color=LIGHTGRAY, font='Arial')
        elif line.startswith('## '):
            p = para(doc, line[3:].strip(), size=14, bold=True, space_after=8,
                     color=DARKGRAY, font='Arial')
            p.paragraph_format.space_before = Pt(16)
            set_bottom_border(p, ACCENT, sz=12, space=4)
        elif line.startswith('### '):
            # акцентный маркер + тёмно-серый заголовок новости
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run('■ ')
            set_font(r, 10, bold=True, color=ACCENT, name='Arial')
            add_runs_with_links(p, line[4:].strip(), 12, True, DARKGRAY, 'Arial')
        elif line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(4)
            add_runs_with_links(p, line[2:].strip(), 12, False, GRAY, 'Times New Roman')
        else:
            para(doc, line.strip())
    doc.save(dst)
    print('Сохранено:', dst)

if __name__ == '__main__':
    build(sys.argv[1], sys.argv[2])