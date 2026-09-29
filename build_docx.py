#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Сборка Дайджеста из markdown-подобного текста в .docx (Times New Roman 12, 1.15)."""
import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_font(run, size=12, bold=False, italic=False, color=None):
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Times New Roman'); rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rPr.append(rFonts)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '24'); rPr.append(sz)  # 12pt
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1'); rPr.append(color)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    new_run.append(rPr)
    t = OxmlElement('w:t'); t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

LINK_RE = re.compile(r'\[([^\]]+)\]\(([^)]+)\)')
BOLD_RE = re.compile(r'\*\*([^*]+)\*\*')

def add_runs_with_links(par, text, base_size=12, base_bold=False):
    """Добавляет текст в параграф, обрабатывая **жирный** и [текст](ссылку)."""
    # Сначала разбиваем по ссылкам, внутри сегментов — по жирному.
    pos = 0
    for m in LINK_RE.finditer(text):
        seg = text[pos:m.start()]
        _add_bold_segments(par, seg, base_size, base_bold)
        add_hyperlink(par, m.group(1), m.group(2))
        pos = m.end()
    _add_bold_segments(par, text[pos:], base_size, base_bold)

def _add_bold_segments(par, text, size, base_bold):
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            r = par.add_run(text[pos:m.start()]); set_font(r, size, base_bold)
        r = par.add_run(m.group(1)); set_font(r, size, True)
        pos = m.end()
    if pos < len(text):
        r = par.add_run(text[pos:]); set_font(r, size, base_bold)

def para(doc, text, size=12, bold=False, space_after=6, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.15
    if align: p.alignment = align
    add_runs_with_links(p, text, size, bold)
    return p

def build(src, dst, title):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2); section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5); section.right_margin = Cm(1.5)
    with open(src, encoding='utf-8') as f:
        lines = f.read().splitlines()
    for line in lines:
        if not line.strip():
            continue
        if line.startswith('# '):
            para(doc, line[2:].strip(), size=16, bold=True, space_after=10, align=1)
        elif line.startswith('## '):
            p = para(doc, line[3:].strip(), size=14, bold=True, space_after=8)
            p.paragraph_format.space_before = Pt(12)
        elif line.startswith('### '):
            p = para(doc, line[4:].strip(), size=12, bold=True, space_after=4)
            p.paragraph_format.space_before = Pt(8)
        elif line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(4)
            add_runs_with_links(p, line[2:].strip(), 12, False)
        elif line.startswith('Дата:'):
            para(doc, line.strip(), size=10, space_after=2)
        else:
            para(doc, line.strip())
    doc.save(dst)
    print('Сохранено:', dst)

if __name__ == '__main__':
    build(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else '')